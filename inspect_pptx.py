from pptx import Presentation
prs = Presentation('Cadence_PII_Shield_Stakeholder_Deck.pptx')
print(f'Total Slides: {len(prs.slides)}')
for i, slide in enumerate(prs.slides):
    title = slide.shapes.title.text if slide.shapes.title else "(no title)"
    print(f'Slide {i+1}: {title}')
