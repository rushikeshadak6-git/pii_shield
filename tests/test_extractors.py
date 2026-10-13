"""
tests/test_extractors.py
------------------------
Integration tests for each extractor on the synthetic documents.
These tests require the synthetic files to have been generated first
(run: python synthetic/generate.py).
"""
import pytest
from pathlib import Path
from extract.router import route
from extract.models import Chunk


# ---------------------------------------------------------------------------
# Portable TrueType font helper (Issue 2)
# ---------------------------------------------------------------------------

_FONT_SEARCH_PATHS = [
    # 1. Repo-bundled (cross-platform; Bitstream Vera / DejaVu, open licence)
    Path(__file__).parent / "assets" / "DejaVuSans.ttf",
    # 2. Common Windows system fonts
    Path(r"C:/Windows/Fonts/arial.ttf"),
    Path(r"C:/Windows/Fonts/verdana.ttf"),
    Path(r"C:/Windows/Fonts/segoeui.ttf"),
    # 3. Common Linux system fonts (DejaVu, Liberation)
    Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    Path("/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"),
    Path("/usr/share/fonts/liberation/LiberationSans-Regular.ttf"),
    Path("/usr/share/fonts/TTF/DejaVuSans.ttf"),
    # 4. macOS
    Path("/Library/Fonts/Arial.ttf"),
    Path("/System/Library/Fonts/Helvetica.ttc"),
]


def _find_font(size: int = 24):
    """Return an ImageFont.FreeTypeFont at *size* pt using the first available
    TrueType font from *_FONT_SEARCH_PATHS*.  If none is found, the calling
    test is skipped with a descriptive message.  Never falls back to the
    default bitmap font (which OCR cannot read reliably).
    """
    from PIL import ImageFont
    for path in _FONT_SEARCH_PATHS:
        if path.exists():
            try:
                return ImageFont.truetype(str(path), size)
            except Exception:
                continue
    pytest.skip(
        "No usable TrueType font found on this system. "
        "Install DejaVu fonts (e.g. `apt-get install fonts-dejavu`) or place "
        "tests/assets/DejaVuSans.ttf in the repo."
    )


@pytest.fixture(scope="module")
def pdf_chunks(config, synthetic_dir):
    pdf = synthetic_dir / "scanned_policy.pdf"
    if not pdf.exists():
        pytest.skip("scanned_policy.pdf not found — run synthetic/generate.py first")
    return route(str(pdf), config)


@pytest.fixture(scope="module")
def docx_chunks(config, synthetic_dir):
    docx = synthetic_dir / "training_guide.docx"
    if not docx.exists():
        pytest.skip("training_guide.docx not found — run synthetic/generate.py first")
    return route(str(docx), config)


@pytest.fixture(scope="module")
def pptx_chunks(config, synthetic_dir):
    pptx = synthetic_dir / "org_pack.pptx"
    if not pptx.exists():
        pytest.skip("org_pack.pptx not found — run synthetic/generate.py first")
    return route(str(pptx), config)


class TestPDFExtractor:
    def test_returns_chunks(self, pdf_chunks):
        assert len(pdf_chunks) > 0

    def test_has_ocr_chunks(self, pdf_chunks):
        """Image-only PDF must have OCR or blocked chunks — no native text chunks."""
        methods = {c.method for c in pdf_chunks}
        # scanned_policy has no text layer, so expect ocr or blocked (not native body)
        assert "native" not in methods or all(
            c.location.startswith("metadata") for c in pdf_chunks if c.method == "native"
        )

    def test_metadata_chunk_present(self, pdf_chunks):
        meta = [c for c in pdf_chunks if "metadata" in c.location]
        assert len(meta) > 0, "PDF metadata chunks should be extracted"

    def test_chunk_has_text_or_blocked(self, pdf_chunks):
        for chunk in pdf_chunks:
            if chunk.is_blocked():
                assert chunk.blocked_reason is not None
            else:
                assert isinstance(chunk.text, str)

    def test_page_locations(self, pdf_chunks):
        page_locs = [c.location for c in pdf_chunks if c.location.startswith("page:")]
        assert len(page_locs) >= 1, "At least one page location expected"

    def test_order_index_monotonic(self, pdf_chunks):
        indices = [c.order_index for c in pdf_chunks]
        assert indices == sorted(indices), "order_index should be monotonically increasing"

    def test_pdf_ocr_caching(self, synthetic_dir, config):
        """Second extraction of a PDF must hit the sha1 page cache."""
        from extract.pdf import extract_pdf, clear_pdf_cache, get_pdf_cache_size

        pdf = synthetic_dir / "scanned_policy.pdf"
        if not pdf.exists():
            pytest.skip("scanned_policy.pdf not found")

        clear_pdf_cache()
        c1 = extract_pdf(str(pdf), config)
        cache_size_after_first = get_pdf_cache_size()
        assert cache_size_after_first > 0, "Pages should be cached"

        c2 = extract_pdf(str(pdf), config)
        assert len(c1) == len(c2)
        assert [c.text for c in c1] == [c.text for c in c2]


class TestDOCXExtractor:
    def test_returns_chunks(self, docx_chunks):
        assert len(docx_chunks) > 0

    def test_has_native_chunks(self, docx_chunks):
        native = [c for c in docx_chunks if c.method == "native"]
        assert len(native) > 0

    def test_metadata_present(self, docx_chunks):
        meta = [c for c in docx_chunks if "metadata" in c.location]
        assert len(meta) > 0

    def test_table_chunks_present(self, docx_chunks):
        tables = [c for c in docx_chunks if "table:" in c.location]
        assert len(tables) > 0, "DOCX has tables — table chunks expected"

    def test_table_chunk_has_pipe_separator(self, docx_chunks):
        tables = [c for c in docx_chunks if "table:" in c.location]
        # At least one table row should have pipe-separated values
        assert any("|" in c.text for c in tables)

    def test_header_chunk_present(self, docx_chunks):
        headers = [c for c in docx_chunks if c.location == "header"]
        assert len(headers) > 0, "DOCX header should be extracted"

    def test_image_chunk_present(self, docx_chunks):
        images = [c for c in docx_chunks if "image:" in c.location]
        assert len(images) > 0, "DOCX embedded PNG should produce an image chunk"

    def test_no_raw_pii_leaks_from_extractor(self, docx_chunks):
        """Extractor itself does NOT redact — this just checks Chunk structure is sane."""
        for c in docx_chunks:
            assert hasattr(c, "text")
            assert hasattr(c, "location")
            assert hasattr(c, "method")

    def test_docx_text_box_extraction(self, tmp_path, config):
        """Text boxes in body (w:txbxContent) must be extracted with location=textbox:N."""
        from docx import Document
        from docx.oxml import OxmlElement
        from extract.docx import extract_docx

        doc = Document()
        p = doc.add_paragraph("Main paragraph text")
        r = p.add_run()
        txbx = OxmlElement("w:txbxContent")
        p_tb = OxmlElement("w:p")
        r_tb = OxmlElement("w:r")
        t_tb = OxmlElement("w:t")
        t_tb.text = "Reviewer: Marcus Vance"
        r_tb.append(t_tb)
        p_tb.append(r_tb)
        txbx.append(p_tb)
        r._r.append(txbx)

        doc_path = tmp_path / "test_tb.docx"
        doc.save(str(doc_path))

        chunks = extract_docx(str(doc_path), config)
        assert any("Marcus Vance" in c.text for c in chunks)
        assert any("textbox:" in c.location for c in chunks)

    def test_docx_nested_tables(self, tmp_path, config):
        """Nested tables within cells must be extracted with nested location path."""
        from docx import Document
        from extract.docx import extract_docx

        doc = Document()
        tbl = doc.add_table(rows=2, cols=2)
        tbl.cell(0, 0).text = "Outer Col 1"
        tbl.cell(0, 1).text = "Outer Col 2"
        tbl.cell(1, 0).text = "Data Cell"

        nested_tbl = tbl.cell(1, 1).add_table(rows=1, cols=2)
        nested_tbl.cell(0, 0).text = "Nested Key"
        nested_tbl.cell(0, 1).text = "Nested Value"

        doc_path = tmp_path / "test_nested.docx"
        doc.save(str(doc_path))

        chunks = extract_docx(str(doc_path), config)
        assert any("Nested Value" in c.text for c in chunks)
        assert any("nested" in c.location for c in chunks)

    def test_docx_tracked_changes_deletion(self, tmp_path, config):
        """Tracked deletions (w:del / w:delText) must be extracted to prevent leaks."""
        from docx import Document
        from docx.oxml import OxmlElement
        from extract.docx import extract_docx

        doc = Document()
        p = doc.add_paragraph("Normal text before deletion. ")
        del_elem = OxmlElement("w:del")
        del_text = OxmlElement("w:delText")
        del_text.text = "Deleted Secret ID: 998-11-2233"
        del_elem.append(del_text)
        p._p.append(del_elem)

        doc_path = tmp_path / "test_del.docx"
        doc.save(str(doc_path))

        chunks = extract_docx(str(doc_path), config)
        assert any("Deleted Secret ID: 998-11-2233" in c.text for c in chunks)
        assert any("tracked_change:" in c.location for c in chunks)

    def test_docx_multi_section_headers(self, tmp_path, config):
        """Headers across multiple sections must all be extracted."""
        from docx import Document
        from extract.docx import extract_docx

        doc = Document()
        sec1 = doc.sections[0]
        sec1.header.paragraphs[0].text = "Header First Section"
        sec2 = doc.add_section()
        sec2.header.is_linked_to_previous = False
        sec2.header.paragraphs[0].text = "Header Second Section"

        doc_path = tmp_path / "test_multisec.docx"
        doc.save(str(doc_path))

        chunks = extract_docx(str(doc_path), config)
        assert any("Header First Section" in c.text for c in chunks)
        assert any("Header Second Section" in c.text for c in chunks)



class TestPPTXExtractor:
    def test_returns_chunks(self, pptx_chunks):
        assert len(pptx_chunks) > 0

    def test_has_native_chunks(self, pptx_chunks):
        native = [c for c in pptx_chunks if c.method == "native"]
        assert len(native) > 0

    def test_slide_locations(self, pptx_chunks):
        slides = [c for c in pptx_chunks if c.location.startswith("slide:")]
        assert len(slides) > 0

    def test_notes_chunks(self, pptx_chunks):
        notes = [c for c in pptx_chunks if "/notes" in c.location]
        assert len(notes) > 0, "Speaker notes should be extracted"

    def test_table_chunks(self, pptx_chunks):
        tables = [c for c in pptx_chunks if "table:" in c.location]
        assert len(tables) > 0, "PPTX tables should be extracted"

    def test_metadata_present(self, pptx_chunks):
        meta = [c for c in pptx_chunks if "metadata" in c.location]
        assert len(meta) > 0


class TestDeskew:
    @pytest.mark.parametrize("angle", [3.0, 5.0, -4.0])
    def test_deskew_rotations(self, angle):
        import difflib
        from PIL import Image, ImageDraw, ImageFont
        from extract.ocr import ocr_image, is_tesseract_available

        if not is_tesseract_available():
            pytest.skip("Tesseract is not available")

        text_lines = [
            "Cadence Financial Services Audit Report",
            "Confidential Employee Compliance Guidelines",
            "Data protection procedures must be strictly followed.",
            "All records herein are synthetic and for verification only.",
            "Authorized Personnel: Operations Oversight Unit"
        ]
        gt_text = " ".join(text_lines)
        img = Image.new("RGB", (1000, 500), (255, 255, 255))
        draw = ImageDraw.Draw(img)
        font = _find_font(24)  # portable helper — skips if no TTF available

        y = 50
        for line in text_lines:
            draw.text((50, y), line, fill=(0, 0, 0), font=font)
            y += 50

        rotated = img.rotate(angle, expand=False, fillcolor=(255, 255, 255))
        ocr_text, conf = ocr_image(rotated)
        clean_ocr = " ".join(ocr_text.split())
        ratio = difflib.SequenceMatcher(None, clean_ocr.lower(), gt_text.lower()).ratio()
        assert ratio >= 0.95, f"Deskew failed at {angle} deg: ratio={ratio:.4f}"


class TestOCRPipeline:
    def test_line_breaks_preserved(self):
        from PIL import Image, ImageDraw, ImageFont
        from extract.ocr import ocr_image, is_tesseract_available

        if not is_tesseract_available():
            pytest.skip("Tesseract is not available")

        lines = [
            "First Line Header",
            "Second Line Department",
            "Third Line Location",
            "Fourth Line Employee",
            "Fifth Line Timestamp",
        ]
        img = Image.new("RGB", (900, 400), (255, 255, 255))
        draw = ImageDraw.Draw(img)
        font = _find_font(24)  # portable helper — skips if no TTF available

        y = 30
        for line in lines:
            draw.text((40, y), line, fill=(0, 0, 0), font=font)
            y += 60

        ocr_text, conf = ocr_image(img)
        assert conf > 50
        extracted_lines = [l.strip() for l in ocr_text.strip().split("\n") if l.strip()]
        assert len(extracted_lines) == 5, f"Expected 5 lines, got {len(extracted_lines)}: {repr(ocr_text)}"

    def test_explicit_blocked_reasons(self, tmp_path, config):
        from PIL import Image
        from unittest.mock import patch
        from extract.router import route
        from extract.ocr import is_tesseract_available

        cfg = config

        # 1. Blank image -> "no text detected"
        blank = Image.new("RGB", (200, 200), (255, 255, 255))
        blank_p = tmp_path / "blank.png"
        blank.save(str(blank_p))
        chunks = route(str(blank_p), cfg)
        assert len(chunks) == 1
        assert chunks[0].is_blocked()
        assert chunks[0].blocked_reason == "no text detected"

        # 2. Mock Tesseract not installed -> "Tesseract not installed"
        with patch("extract.router.is_tesseract_available", return_value=False):
            chunks = route(str(blank_p), cfg)
            assert len(chunks) == 1
            assert chunks[0].is_blocked()
            assert chunks[0].blocked_reason == "Tesseract not installed"

        # 3. Sub-threshold OCR text -> low_confidence=True and blocked_reason has "below threshold"
        with patch("extract.router.ocr_image", return_value=("Some faint noisy text", 25.0)):
            chunks = route(str(blank_p), cfg)
            assert len(chunks) == 1
            assert chunks[0].is_blocked()
            assert chunks[0].low_confidence is True
            assert "below threshold" in chunks[0].blocked_reason
            assert chunks[0].text == "Some faint noisy text"
