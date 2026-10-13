"""
synthetic_holdout/generate_holdout.py
-------------------------------------
Generate an independent, hold-out synthetic test set with seed 99 (never used for training/tuning).

Features:
  - New templates and organization names: "Vanguard Logistics", "Apex Health"
  - Table structures with columns: Users, Assignee, Client Contact
  - Scanned PDF with a +3.0° rotation on pages (tests OCR deskew fix)
  - Rich DOCX with first-page header, text box, tracked changes, nested tables
  - Generates ground_truth/*.json and ground_truth/*_text.txt
"""
from __future__ import annotations

import io
import json
import os
import random
import sys
from pathlib import Path
from typing import Dict, List

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

import fitz  # PyMuPDF
from faker import Faker
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor, Inches
from pptx import Presentation
from pptx.util import Inches as PInches, Pt as PPt

fake = Faker("en_US")
fake.seed_instance(99)
random.seed(99)

OUT_DIR = Path(__file__).resolve().parent
GT_DIR = OUT_DIR / "ground_truth"
GT_DIR.mkdir(parents=True, exist_ok=True)


def _fake_ssn() -> str:
    while True:
        area = random.randint(1, 899)
        if area == 666:
            continue
        group = random.randint(1, 99)
        serial = random.randint(1, 9999)
        if area != 0 and group != 0 and serial != 0:
            return f"{area:03d}-{group:02d}-{serial:04d}"


def _fake_employee_id() -> str:
    return f"EMP-{random.randint(10000, 99999)}"


def _fake_ip() -> str:
    return f"10.{random.randint(1,254)}.{random.randint(1,254)}.{random.randint(1,254)}"


def _gt_entity(etype: str, text: str, location: str, tricky: bool = False) -> Dict:
    return {"type": etype, "text": text, "location": location, "tricky": tricky}


def _save_gt(filename: str, entities: List[Dict], doc_name: str) -> None:
    path = GT_DIR / filename
    with open(path, "w", encoding="utf-8") as fh:
        json.dump({"document": doc_name, "entities": entities}, fh, indent=2)
    print(f"  Ground truth → {path}")


def _save_gt_text(filename: str, text: str) -> None:
    path = GT_DIR / filename
    path.write_text(text, encoding="utf-8")
    print(f"  GT text      → {path}")


# ---------------------------------------------------------------------------
# 1. Scanned PDF (with +3.0° rotation)
# ---------------------------------------------------------------------------

def _render_page_image(text_lines: List[str], rotate_deg: float = 3.0, noise_level: int = 4) -> Image.Image:
    W, H = 794, 1123
    img = Image.new("RGB", (W, H), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)

    try:
        font = ImageFont.truetype("arial.ttf", 16)
    except OSError:
        font = ImageFont.load_default()

    y = 60
    for line in text_lines:
        draw.text((60, y), line, fill=(10, 10, 10), font=font)
        y += 24
        if y > H - 60:
            break

    if abs(rotate_deg) > 0.01:
        img = img.rotate(rotate_deg, expand=False, fillcolor=(255, 255, 255))

    if noise_level > 0:
        import numpy as np
        arr = np.array(img, dtype=np.int16)
        noise = np.random.randint(-noise_level * 2, noise_level * 2, arr.shape, dtype=np.int16)
        arr = np.clip(arr + noise, 0, 255).astype(np.uint8)
        img = Image.fromarray(arr)

    return img


def generate_scanned_pdf() -> None:
    doc_name = "scanned_policy.pdf"
    out_path = OUT_DIR / doc_name
    print(f"\n[1/3] Generating {doc_name} with +3.0° skew…")

    entities: List[Dict] = []
    gt_text_lines: List[str] = []

    author = fake.name()
    names = [fake.name() for _ in range(4)]
    emails = [fake.email() for _ in range(3)]
    phones = [fake.phone_number() for _ in range(3)]
    ssn = _fake_ssn()
    emp_id = _fake_employee_id()
    ip = _fake_ip()

    entities.append(_gt_entity("PERSON", author, "metadata/author"))

    pages_content: List[List[str]] = []

    # Page 1
    p1 = [
        "VANGUARD LOGISTICS CORP",
        "GLOBAL SUPPLY CHAIN SECURITY POLICY",
        f"Author: {author}",
        "Version 4.0  |  Effective: 2025-03-01",
        "",
        "Confidential document. Unauthorized distribution prohibited.",
        "Inquiries: security@vanguard-logistics.com",
    ]
    pages_content.append(p1)
    entities.append(_gt_entity("PERSON", author, "page:1"))
    entities.append(_gt_entity("EMAIL", "security@vanguard-logistics.com", "page:1"))
    gt_text_lines.extend(p1)

    # Page 2
    p2 = [
        "SECTION 2: PERSONNEL CREDENTIALS",
        "",
        f"Employee Name: {names[0]}",
        f"Employee ID: {emp_id}",
        f"Social Security: {ssn}",
        f"Email: {emails[0]}",
        f"Telephone: {phones[0]}",
        f"Workstation IP: {ip}",
        "",
        "Fleet tracking authorization is granted under ISO-27001.",
    ]
    pages_content.append(p2)
    entities.append(_gt_entity("PERSON", names[0], "page:2"))
    entities.append(_gt_entity("EMPLOYEE_ID", emp_id, "page:2"))
    entities.append(_gt_entity("SSN", ssn, "page:2"))
    entities.append(_gt_entity("EMAIL", emails[0], "page:2"))
    entities.append(_gt_entity("PHONE", phones[0], "page:2"))
    entities.append(_gt_entity("IP_ADDRESS", ip, "page:2"))
    gt_text_lines.extend(p2)

    # Page 3
    p3 = [
        "SECTION 3: REGIONAL DIRECTORS",
        "",
        f"Operations Director: {names[1]}",
        f"Direct Phone: {phones[1]}",
        f"Email: {emails[1]}",
        "",
        "Terminal coordinates and manifest routing are monitored 24/7.",
    ]
    pages_content.append(p3)
    entities.append(_gt_entity("PERSON", names[1], "page:3"))
    entities.append(_gt_entity("PHONE", phones[1], "page:3"))
    entities.append(_gt_entity("EMAIL", emails[1], "page:3"))
    gt_text_lines.extend(p3)

    # Page 4
    p4 = [
        "SECTION 4: COMPLIANCE SIGN-OFF",
        "",
        f"Reviewer: {names[2]}",
        f"Contact: {emails[2]}",
        f"Mobile: {phones[2]}",
        "",
        f"Final Approval by Document Owner: {names[3]}",
    ]
    pages_content.append(p4)
    entities.append(_gt_entity("PERSON", names[2], "page:4"))
    entities.append(_gt_entity("EMAIL", emails[2], "page:4"))
    entities.append(_gt_entity("PHONE", phones[2], "page:4"))
    entities.append(_gt_entity("PERSON", names[3], "page:4"))
    gt_text_lines.extend(p4)

    # Render PDF with +3.0 degree rotation (tests deskew)
    pdf_doc = fitz.open()
    rotations = [3.0, 3.0, 3.0, 3.0]
    for lines, rot in zip(pages_content, rotations):
        img = _render_page_image(lines, rotate_deg=rot)
        buf = io.BytesIO()
        img.save(buf, format="PNG")
        buf.seek(0)
        page = pdf_doc.new_page(width=595, height=842)
        page.insert_image(fitz.Rect(0, 0, 595, 842), stream=buf.read())

    pdf_doc.set_metadata({
        "author": author,
        "title": "Global Supply Chain Security Policy",
        "subject": "Vanguard Logistics Policy",
    })
    pdf_doc.save(str(out_path))
    pdf_doc.close()
    print(f"  Written      → {out_path}")

    _save_gt("scanned_policy.json", entities, doc_name)
    _save_gt_text("scanned_policy_text.txt", "\n".join(gt_text_lines))


# ---------------------------------------------------------------------------
# 2. Rich DOCX (Nested tables, text box, tracked changes, first-page header)
# ---------------------------------------------------------------------------

def generate_docx() -> None:
    doc_name = "training_guide.docx"
    out_path = OUT_DIR / doc_name
    print(f"\n[2/3] Generating {doc_name}…")

    entities: List[Dict] = []
    gt_text_lines: List[str] = []

    author = fake.name()
    doc = Document()

    # Core properties
    doc.core_properties.author = author
    doc.core_properties.title = "Apex Health Compliance Standard"
    entities.append(_gt_entity("PERSON", author, "metadata/author"))

    # First section header
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True
    f_header_text = f"Apex Health confidential — Custodian: {author}"
    sec.first_page_header.paragraphs[0].text = f_header_text
    entities.append(_gt_entity("PERSON", author, "first_page_header"))
    gt_text_lines.append(f_header_text)

    # Body Title
    doc.add_heading("Apex Health Patient Safety & Operations Guide", 0)
    gt_text_lines.append("Apex Health Patient Safety & Operations Guide")

    # Table with columns: Users | Assignee | Client Contact
    doc.add_heading("1. Operational Team Roster", 1)
    gt_text_lines.append("1. Operational Team Roster")

    table1 = doc.add_table(rows=1, cols=3)
    hdr_cells = table1.rows[0].cells
    hdr_cells[0].text = "Users"
    hdr_cells[1].text = "Assignee"
    hdr_cells[2].text = "Client Contact"
    gt_text_lines.append("Users | Assignee | Client Contact")

    # Initials style names and full names
    user1 = "TJ Harlan"
    assignee1 = fake.name()
    contact1 = fake.name()

    row1 = table1.add_row().cells
    row1[0].text = user1
    row1[1].text = assignee1
    row1[2].text = contact1
    entities.append(_gt_entity("PERSON", user1, "table:1"))
    entities.append(_gt_entity("PERSON", assignee1, "table:1"))
    entities.append(_gt_entity("PERSON", contact1, "table:1"))
    gt_text_lines.append(f"{user1} | {assignee1} | {contact1}")

    user2 = "PQ Lindqvist"
    assignee2 = fake.name()
    contact2 = fake.name()

    row2 = table1.add_row().cells
    row2[0].text = user2
    row2[1].text = assignee2
    row2[2].text = contact2
    entities.append(_gt_entity("PERSON", user2, "table:1"))
    entities.append(_gt_entity("PERSON", assignee2, "table:1"))
    entities.append(_gt_entity("PERSON", contact2, "table:1"))
    gt_text_lines.append(f"{user2} | {assignee2} | {contact2}")

    # Text Box (w:txbxContent)
    tb_name = fake.name()
    tb_email = fake.email()
    p_tb = doc.add_paragraph("Emergency escalation contacts are listed below:")
    gt_text_lines.append("Emergency escalation contacts are listed below:")
    r_tb = p_tb.add_run()
    txbx_elem = OxmlElement("w:txbxContent")
    p_inner = OxmlElement("w:p")
    r_inner = OxmlElement("w:r")
    t_inner = OxmlElement("w:t")
    tb_content = f"Emergency Contact: {tb_name} | Email: {tb_email}"
    t_inner.text = tb_content
    r_inner.append(t_inner)
    p_inner.append(r_inner)
    txbx_elem.append(p_inner)
    r_tb._r.append(txbx_elem)
    entities.append(_gt_entity("PERSON", tb_name, "textbox:1"))
    entities.append(_gt_entity("EMAIL", tb_email, "textbox:1"))
    gt_text_lines.append(tb_content)

    # Tracked change deletion
    del_ssn = _fake_ssn()
    p_del = doc.add_paragraph("Historical identity tracking: ")
    del_elem = OxmlElement("w:del")
    del_text = OxmlElement("w:delText")
    del_content = f"Deleted Social Security Number: {del_ssn}"
    del_text.text = del_content
    del_elem.append(del_text)
    p_del._p.append(del_elem)
    entities.append(_gt_entity("SSN", del_ssn, "tracked_change:deleted/1"))
    gt_text_lines.append(f"Historical identity tracking: {del_content}")

    # Nested Table
    doc.add_heading("2. Clinical Departments", 1)
    gt_text_lines.append("2. Clinical Departments")
    outer_table = doc.add_table(rows=2, cols=2)
    outer_table.cell(0, 0).text = "Department"
    outer_table.cell(0, 1).text = "Unit Lead"
    lead1 = fake.name()
    outer_table.cell(1, 0).text = "Cardiology"
    outer_table.cell(1, 1).text = lead1
    entities.append(_gt_entity("PERSON", lead1, "table:2"))
    gt_text_lines.append(f"Department | Unit Lead")
    gt_text_lines.append(f"Cardiology | {lead1}")

    # Add nested table in cell(1, 0)
    nested_table = outer_table.cell(1, 0).add_table(rows=2, cols=2)
    nested_table.cell(0, 0).text = "Specialist"
    nested_table.cell(0, 1).text = "Direct Line"
    spec_name = fake.name()
    spec_phone = fake.phone_number()
    nested_table.cell(1, 0).text = spec_name
    nested_table.cell(1, 1).text = spec_phone
    entities.append(_gt_entity("PERSON", spec_name, "table:2/nested:1"))
    entities.append(_gt_entity("PHONE", spec_phone, "table:2/nested:1"))
    gt_text_lines.append(f"Specialist | Direct Line")
    gt_text_lines.append(f"{spec_name} | {spec_phone}")

    # Additional filler
    for _ in range(5):
        para = fake.paragraph(nb_sentences=4)
        doc.add_paragraph(para)
        gt_text_lines.append(para)

    doc.save(str(out_path))
    print(f"  Written      → {out_path}")

    _save_gt("training_guide.json", entities, doc_name)
    _save_gt_text("training_guide_text.txt", "\n".join(gt_text_lines))


# ---------------------------------------------------------------------------
# 3. PPTX (Org Presentation with directory, notes)
# ---------------------------------------------------------------------------

def generate_pptx() -> None:
    doc_name = "org_pack.pptx"
    out_path = OUT_DIR / doc_name
    print(f"\n[3/3] Generating {doc_name}…")

    entities: List[Dict] = []
    gt_text_lines: List[str] = []

    prs = Presentation()
    prs.slide_width = PInches(13.33)
    prs.slide_height = PInches(7.5)

    # Slide 1: Title
    slide_layout = prs.slide_layouts[0]
    slide1 = prs.slides.add_slide(slide_layout)
    slide1.shapes.title.text = "Vanguard Logistics & Apex Health"
    slide1.placeholders[1].text = "Joint Strategic Operations 2025"
    gt_text_lines.append("Vanguard Logistics & Apex Health")
    gt_text_lines.append("Joint Strategic Operations 2025")

    # Slide 2: Management Directory Table
    slide2 = prs.slides.add_slide(prs.slide_layouts[5])
    slide2.shapes.title.text = "Executive Directory"
    gt_text_lines.append("Executive Directory")

    rows, cols = 4, 3
    table_shape = slide2.shapes.add_table(rows, cols, PInches(1), PInches(2), PInches(11), PInches(3))
    table = table_shape.table

    table.cell(0, 0).text = "Assignee"
    table.cell(0, 1).text = "Email"
    table.cell(0, 2).text = "Phone"
    gt_text_lines.append("Assignee | Email | Phone")

    for i in range(1, 4):
        nm = fake.name()
        em = fake.email()
        ph = fake.phone_number()
        table.cell(i, 0).text = nm
        table.cell(i, 1).text = em
        table.cell(i, 2).text = ph
        entities.append(_gt_entity("PERSON", nm, f"slide:2/table:1"))
        entities.append(_gt_entity("EMAIL", em, f"slide:2/table:1"))
        entities.append(_gt_entity("PHONE", ph, f"slide:2/table:1"))
        gt_text_lines.append(f"{nm} | {em} | {ph}")

    # Notes on Slide 2
    speaker = fake.name()
    notes_slide = slide2.notes_slide
    notes_slide.notes_text_frame.text = f"Session moderated by {speaker}. Keep directory confidential."
    entities.append(_gt_entity("PERSON", speaker, "slide:2/notes"))
    gt_text_lines.append(f"Session moderated by {speaker}. Keep directory confidential.")

    prs.save(str(out_path))
    print(f"  Written      → {out_path}")

    _save_gt("org_pack.json", entities, doc_name)
    _save_gt_text("org_pack_text.txt", "\n".join(gt_text_lines))


def main():
    print("=" * 60)
    print("GENERATING SYNTHETIC HOLDOUT TEST SET (SEED 99)")
    print("=" * 60)
    generate_scanned_pdf()
    generate_docx()
    generate_pptx()
    print("\nHold-out generation complete.")


if __name__ == "__main__":
    main()
