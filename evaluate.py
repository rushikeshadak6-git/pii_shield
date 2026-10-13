"""
evaluate.py
-----------
End-to-end evaluation runner:
Processes all synthetic documents, measures detection performance against
ground truth, verifies zero leakage, measures throughput, and supports
a 6-stage ablation study (--ablation).
"""
from __future__ import annotations

import argparse
import json
import logging
import sys
import time
from pathlib import Path
from typing import Dict, List, Any

import yaml

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from extract.router import route
from detect.merge import detect_chunk, Entity
from redact.mapping import PlaceholderMapping
from redact.redactor import redact_chunks
from verify.leakage import check_leakage
from verify.metrics import run_evaluation
from report.report import build_file_summary, print_summary_table

log = logging.getLogger("evaluate")
logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")

ROOT = Path(__file__).resolve().parent


def _get_doc_page_count(doc_path: Path) -> int:
    """Estimate or count pages/slides in a document for throughput tracking."""
    ext = doc_path.suffix.lower()
    try:
        if ext == ".pdf":
            import fitz
            doc = fitz.open(doc_path)
            return len(doc)
        elif ext == ".pptx":
            from pptx import Presentation
            prs = Presentation(doc_path)
            return len(prs.slides)
        elif ext == ".docx":
            from docx import Document
            doc = Document(doc_path)
            return max(1, len(doc.paragraphs) // 15)
    except Exception:
        pass
    return 1


def run_ablation_study(
    input_dir: Path,
    gt_dir: Path,
    config: dict,
    doc_files: List[Path],
    output_path: Path,
) -> None:
    """
    Run 6 ablation stages reporting:
    1. rules only
    2. rules + NER
    3. + context rules
    4. + OCR repair (Item 2)
    5. + suspect-line redaction (Item 1)
    6. + alias propagation (Item 3)
    Output a table per dataset (precision, recall, F1, and counts of SUSPECT_OCR and alias redactions).
    """
    print("\n" + "=" * 80)
    print(f"PII SHIELD — 6-STAGE ABLATION STUDY ({input_dir.name})")
    print("=" * 80)

    # 1. Pre-extract all document chunks once
    print("Pre-extracting document chunks...")
    doc_chunks: Dict[str, List] = {}
    extracted_texts: Dict[str, str] = {}
    for doc_path in doc_files:
        chunks = route(str(doc_path), config)
        doc_chunks[doc_path.name] = chunks
        extracted_texts[doc_path.name] = "\n".join(c.text for c in chunks if not c.is_blocked())
    print(f"Extraction complete ({len(doc_files)} files).\n")

    stages = [
        (
            "1. rules only",
            {"use_ner": False, "use_context": False, "use_repair": False, "use_suspect": False, "use_alias": False},
        ),
        (
            "2. rules + NER",
            {"use_ner": True, "use_context": False, "use_repair": False, "use_suspect": False, "use_alias": False},
        ),
        (
            "3. + context rules",
            {"use_ner": True, "use_context": True, "use_repair": False, "use_suspect": False, "use_alias": False},
        ),
        (
            "4. + OCR repair (Item 2)",
            {"use_ner": True, "use_context": True, "use_repair": True, "use_suspect": False, "use_alias": False},
        ),
        (
            "5. + suspect-line (Item 1)",
            {"use_ner": True, "use_context": True, "use_repair": True, "use_suspect": True, "use_alias": False},
        ),
        (
            "6. + alias propagation (Item 3)",
            {"use_ner": True, "use_context": True, "use_repair": True, "use_suspect": True, "use_alias": True},
        ),
    ]

    ablation_results = []

    for stage_name, flags in stages:
        print(f"Running ablation: {stage_name}...")
        stage_pred: List[Dict] = []
        suspect_count = 0
        alias_count = 0

        for doc_path in doc_files:
            chunks = doc_chunks[doc_path.name]
            doc_entities: List[Entity] = []
            for chunk in chunks:
                if not chunk.is_blocked():
                    ents = detect_chunk(
                        chunk,
                        config,
                        use_ner=flags["use_ner"],
                        use_context=flags["use_context"],
                        use_repair=flags["use_repair"],
                        use_suspect=flags["use_suspect"],
                    )
                    doc_entities.extend(ents)

            # Optional stage: alias propagation
            if flags["use_alias"]:
                try:
                    from detect.propagate import propagate_name_aliases
                    person_ents = [e for e in doc_entities if e.entity_type == "PERSON"]
                    chunk_texts = [c.text for c in chunks if not c.is_blocked()]
                    chunk_locs  = [c.location for c in chunks if not c.is_blocked()]
                    alias_ents = propagate_name_aliases(
                        chunk_texts, person_ents, doc_entities, config,
                        file=str(doc_path), locations=chunk_locs,
                    )
                    doc_entities.extend(alias_ents)
                except Exception:
                    pass

            for e in doc_entities:
                if e.entity_type == "SUSPECT_OCR":
                    suspect_count += 1
                if "alias_propagation" in getattr(e, "methods", []):
                    alias_count += 1

                stage_pred.append({
                    "type": e.entity_type,
                    "text": e.raw_value,
                    "location": e.location,
                    "document": doc_path.name,
                })

        eval_res = run_evaluation(str(gt_dir), stage_pred, extracted_texts)
        overall = eval_res.get("entity_metrics", {}).get("OVERALL", {})
        p = overall.get("precision", 0.0)
        r = overall.get("recall", 0.0)
        f1 = overall.get("f1", 0.0)

        ablation_results.append({
            "stage": stage_name,
            "precision": p,
            "recall": r,
            "f1": f1,
            "suspect_ocr_count": suspect_count,
            "alias_redactions_count": alias_count,
            "tp": overall.get("true_positives", 0),
            "fp": overall.get("false_positives", 0),
            "fn": overall.get("false_negatives", 0),
        })

    # Print Formatted Ablation Table
    print("\n" + "=" * 90)
    print(f"ABLATION STUDY SUMMARY TABLE ({input_dir.name})")
    print("=" * 90)
    print(f"{'Stage':<35} {'Prec':>8} {'Recall':>8} {'F1':>8} {'SUSPECT_OCR':>13} {'Aliases':>9}")
    print("-" * 90)
    for row in ablation_results:
        print(
            f"{row['stage']:<35} "
            f"{row['precision']:>8.4f} "
            f"{row['recall']:>8.4f} "
            f"{row['f1']:>8.4f} "
            f"{row['suspect_ocr_count']:>13d} "
            f"{row['alias_redactions_count']:>9d}"
        )
    print("=" * 90)

    # Save ablation results
    ablation_out = input_dir.parent / f"ablation_{input_dir.name}.json"
    with open(ablation_out, "w", encoding="utf-8") as fh:
        json.dump(ablation_results, fh, indent=2)
    print(f"Ablation results saved to: {ablation_out}")


def main():
    parser = argparse.ArgumentParser(description="PII Shield Evaluation")
    parser.add_argument("--input", default="synthetic", help="Input directory containing documents and ground_truth/")
    parser.add_argument("--baseline", default=None, help="Path to baseline json file for comparison")
    parser.add_argument("--output", default="evaluation_results.json", help="Path to save evaluation results json")
    parser.add_argument("--ablation", action="store_true", help="Run 6-stage ablation experiment and report table")
    args = parser.parse_args()

    input_dir = Path(args.input)
    if not input_dir.is_absolute():
        input_dir = ROOT / input_dir
    gt_dir = input_dir / "ground_truth"

    with open(ROOT / "config.yaml", "r", encoding="utf-8") as fh:
        config = yaml.safe_load(fh)

    doc_files = sorted(
        [p for p in input_dir.iterdir() if p.is_file() and p.suffix.lower() in (".docx", ".pptx", ".pdf")]
    )

    if args.ablation:
        run_ablation_study(input_dir, gt_dir, config, doc_files, Path(args.output))
        return

    print("=" * 70)
    print(f"PII SHIELD — COMPREHENSIVE PIPELINE EVALUATION ({input_dir.name})")
    print("=" * 70)

    all_predicted_entities: List[Dict] = []
    extracted_texts: Dict[str, str] = {}
    file_summaries = []
    throughput_stats = []

    for doc_path in doc_files:
        print(f"\nProcessing: {doc_path.name}...")
        t_start = time.perf_counter()

        chunks = route(str(doc_path), config)

        # Reconstruct extracted text
        plain_text = "\n".join(c.text for c in chunks if not c.is_blocked())
        extracted_texts[doc_path.name] = plain_text

        # Detect PII
        entities: List[Entity] = []
        for chunk in chunks:
            if not chunk.is_blocked():
                ents = detect_chunk(chunk, config)
                entities.extend(ents)

        # Item 3: propagate name aliases across document chunks
        try:
            from detect.propagate import propagate_name_aliases
            person_ents = [e for e in entities if e.entity_type == "PERSON"]
            chunk_texts = [c.text for c in chunks if not c.is_blocked()]
            chunk_locs  = [c.location for c in chunks if not c.is_blocked()]
            alias_ents = propagate_name_aliases(
                chunk_texts, person_ents, entities, config,
                file=str(doc_path), locations=chunk_locs,
            )
            entities.extend(alias_ents)
        except Exception:
            pass  # Fail-safe

        # Redact
        out_dir = ROOT / "output_eval"
        mapping = PlaceholderMapping(source_file=str(doc_path), output_dir=str(out_dir))
        redacted = redact_chunks(chunks, entities, mapping, config)

        # Leakage check with known-value propagation
        leak_res = check_leakage(redacted, config, known_values=mapping)

        # File summary
        summary = build_file_summary(
            str(doc_path),
            chunks,
            entities,
            leak_result=leak_res,
            structure_score=None,
        )
        file_summaries.append(summary)

        # Collect predicted entities for metrics
        for e in entities:
            all_predicted_entities.append({
                "type": e.entity_type,
                "text": e.raw_value,
                "location": e.location,
                "document": doc_path.name,
            })

        t_elapsed = time.perf_counter() - t_start
        page_count = _get_doc_page_count(doc_path)
        pages_per_min = (page_count / t_elapsed) * 60.0 if t_elapsed > 0 else 0.0

        throughput_stats.append({
            "file": doc_path.name,
            "seconds": t_elapsed,
            "pages": page_count,
            "ppm": pages_per_min,
        })

        print(f"  Chunks extracted : {len(chunks)} (Blocked: {summary['blocked_chunks']})")
        print(f"  PII detected     : {len(entities)}")
        print(f"  Leakage status   : {'CLEAN (Safe)' if leak_res.downstream_safe else 'LEAK DETECTED'}")
        print(f"  Throughput       : {t_elapsed:.2f}s ({pages_per_min:.1f} pages/min)")

    print("\n" + "=" * 70)
    print("AUDIT SUMMARY REPORT TABLE")
    print("=" * 70)
    print_summary_table(file_summaries)

    print("\n" + "=" * 70)
    print("THROUGHPUT PERFORMANCE")
    print("=" * 70)
    print(f"{'File':<28} {'Time (s)':>10} {'Pages/Slides':>14} {'Speed (PPM)':>14}")
    print("-" * 70)
    for tp in throughput_stats:
        print(f"{tp['file']:<28} {tp['seconds']:>10.2f} {tp['pages']:>14d} {tp['ppm']:>14.1f}")
    print("-" * 70)
    total_time = sum(t["seconds"] for t in throughput_stats)
    total_pages = sum(t["pages"] for t in throughput_stats)
    avg_ppm = (total_pages / total_time) * 60.0 if total_time > 0 else 0.0
    print(f"{'TOTAL / AVERAGE':<28} {total_time:>10.2f} {total_pages:>14d} {avg_ppm:>14.1f}")
    print("=" * 70)

    print("\n" + "=" * 70)
    print("BENCHMARK METRICS (VS GROUND TRUTH)")
    print("=" * 70)
    eval_results = run_evaluation(str(gt_dir), all_predicted_entities, extracted_texts)
    eval_results["throughput"] = throughput_stats

    eval_out = Path(args.output)
    if not eval_out.is_absolute():
        eval_out = ROOT / eval_out
    with open(eval_out, "w", encoding="utf-8") as fh:
        json.dump(eval_results, fh, indent=2)
    print(f"\nSaved evaluation metrics to {eval_out}")

    # Baseline comparison if provided
    if args.baseline:
        b_path = Path(args.baseline)
        if not b_path.is_absolute():
            b_path = ROOT / b_path
        if b_path.exists():
            with open(b_path, "r", encoding="utf-8") as fh:
                baseline_data = json.load(fh)
            print("\n" + "=" * 70)
            print("COMPARISON VS BASELINE")
            print("=" * 70)
            b_overall = baseline_data.get("entity_metrics", {}).get("OVERALL", {})
            c_overall = eval_results.get("entity_metrics", {}).get("OVERALL", {})
            b_rec = b_overall.get("recall", 0.0)
            c_rec = c_overall.get("recall", 0.0)
            b_prec = b_overall.get("precision", 0.0)
            c_prec = c_overall.get("precision", 0.0)
            b_f1 = b_overall.get("f1", 0.0)
            c_f1 = c_overall.get("f1", 0.0)
            print(f"Overall Recall   : {b_rec:.4f} → {c_rec:.4f} ({c_rec - b_rec:+.4f})")
            print(f"Overall Precision: {b_prec:.4f} → {c_prec:.4f} ({c_prec - b_prec:+.4f})")
            print(f"Overall F1       : {b_f1:.4f} → {c_f1:.4f} ({c_f1 - b_f1:+.4f})")


if __name__ == "__main__":
    main()
